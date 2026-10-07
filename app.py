import streamlit as st
import openpyxl
import subprocess
import tempfile
import os
import shutil
import json
import urllib.request
import urllib.error
from datetime import datetime, timezone, timedelta

try:
    from pypdf import PdfMerger
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

st.set_page_config(
    page_title="ระบบแปลง Excel เป็น PDF",
    page_icon="📄",
    layout="centered"
)

st.title("📄 ระบบแปลงรายงาน Excel เป็น PDF")
st.markdown("อัปโหลดไฟล์ **Excel (.xlsx)** เพื่อรับไฟล์ **PDF** ที่จัดหน้าและตัดเฉพาะข้อมูลส่วนที่จำเป็นให้อัตโนมัติ")

def extract_summary_data(file_path):
    """ดึงยอดสรุปจากทั้ง 2 ชีท โดยใช้ data_only=True เพื่ออ่านผลลัพธ์ตัวเลขจริง"""
    summary = {
        "ws1_update": "",
        "ws1_billed": None,
        "ws1_pending": None,
        "ws1_total": None,
        "ws2_update": "",
        "ws2_total": None
    }
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        # 1. ชีท: ยอดค้างรับลูกค้า
        if "ยอดค้างรับลูกค้า" in wb.sheetnames:
            ws = wb["ยอดค้างรับลูกค้า"]
            for r in range(1, ws.max_row + 1):
                for c in range(1, 10):
                    val = str(ws.cell(r, c).value or "").strip()
                    if "Update" in val and not summary["ws1_update"]:
                        summary["ws1_update"] = val
                    elif "วางบิลแล้ว" in val and summary["ws1_billed"] is None:
                        for col_idx in range(c + 1, min(c + 6, ws.max_column + 1)):
                            v = ws.cell(r, col_idx).value
                            if isinstance(v, (int, float)) and v > 0:
                                summary["ws1_billed"] = v
                                break
                    elif "รอวางบิล" in val and summary["ws1_pending"] is None:
                        for col_idx in range(c + 1, min(c + 6, ws.max_column + 1)):
                            v = ws.cell(r, col_idx).value
                            if isinstance(v, (int, float)):
                                summary["ws1_pending"] = v
                                break
                    elif "ยอดรวมวางบิลทั้งหมด" in val and summary["ws1_total"] is None:
                        for col_idx in range(c + 1, min(c + 6, ws.max_column + 1)):
                            v = ws.cell(r, col_idx).value
                            if isinstance(v, (int, float)) and v > 0:
                                summary["ws1_total"] = v
                                break

        # 2. ชีท: ยอดค้างจ่ายหัวลาก
        if "ยอดค้างจ่ายหัวลาก" in wb.sheetnames:
            ws = wb["ยอดค้างจ่ายหัวลาก"]
            for r in range(1, ws.max_row + 1):
                for c in range(1, 10):
                    val = str(ws.cell(r, c).value or "").strip()
                    if "Update" in val and not summary["ws2_update"]:
                        summary["ws2_update"] = val
                        # ตรวจหาแถวสรุปยอดรวมถัดไป 1-4 แถว
                        for next_r in range(r, min(r + 5, ws.max_row + 1)):
                            for next_c in range(1, ws.max_column + 1):
                                cell_str = str(ws.cell(next_r, next_c).value or "").strip()
                                if "ยอดรวม" in cell_str:
                                    for col_idx in range(next_c + 1, min(next_c + 6, ws.max_column + 1)):
                                        v = ws.cell(next_r, col_idx).value
                                        if isinstance(v, (int, float)) and v > 0:
                                            summary["ws2_total"] = v
                                            break
                            if summary["ws2_total"] is not None:
                                break
    except Exception as e:
        print(f"Error extracting summary: {e}")

    return summary

def send_line_notification(file_name, summary):
    """ส่งข้อความแจ้งเตือนเข้า LINE พร้อมตัวเลขสรุปยอด"""
    try:
        token = st.secrets.get("LINE_CHANNEL_ACCESS_TOKEN")
        user_id = st.secrets.get("LINE_USER_ID")

        if not token or not user_id:
            return False, "ยังไม่ได้ตั้งค่า LINE_CHANNEL_ACCESS_TOKEN หรือ LINE_USER_ID ใน Streamlit Secrets"

        bkk_tz = timezone(timedelta(hours=7))
        now_str = datetime.now(bkk_tz).strftime("%d/%m/%Y %H:%M:%S")

        msg_lines = [
            "📄 แปลงไฟล์ Excel เป็น PDF สำเร็จ!",
            f"📁 ไฟล์: {file_name}",
            f"⏰ เวลา: {now_str}",
            ""
        ]

        # ข้อมูลยอดค้างรับลูกค้า
        ws1_up = f" ({summary['ws1_update']})" if summary.get("ws1_update") else ""
        msg_lines.append(f"📊 [ยอดค้างรับลูกค้า]{ws1_up}")
        if summary.get("ws1_billed") is not None:
            msg_lines.append(f"• วางบิลแล้ว: {summary['ws1_billed']:,.2f} บาท")
        if summary.get("ws1_pending") is not None:
            msg_lines.append(f"• รอวางบิล: {summary['ws1_pending']:,.2f} บาท")
        if summary.get("ws1_total") is not None:
            msg_lines.append(f"• รวมทั้งหมด: {summary['ws1_total']:,.2f} บาท")
        elif summary.get("ws1_billed") is None and summary.get("ws1_pending") is None and summary.get("ws1_total") is None:
            msg_lines.append("• แปลงเอกสารเรียบร้อย")

        msg_lines.append("")

        # ข้อมูลยอดค้างจ่ายหัวลาก
        ws2_up = f" ({summary['ws2_update']})" if summary.get("ws2_update") else ""
        msg_lines.append(f"🚛 [ยอดค้างจ่ายหัวลาก]{ws2_up}")
        if summary.get("ws2_total") is not None:
            msg_lines.append(f"• ยอดรวมค้างจ่าย: {summary['ws2_total']:,.2f} บาท")
        else:
            msg_lines.append("• แปลงเอกสารเรียบร้อย")

        full_msg = "\n".join(msg_lines)

        url = "https://api.line.me/v2/bot/message/push"
        headers = {
            "Content-Type": "application/json",
            "Authorization": f"Bearer {token}"
        }
        payload = {
            "to": user_id,
            "messages": [{"type": "text", "text": full_msg}]
        }

        req = urllib.request.Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            headers=headers
        )
        with urllib.request.urlopen(req, timeout=10) as res:
            return True, "ส่งข้อความสำเร็จ"
    except urllib.error.HTTPError as e:
        err_body = e.read().decode("utf-8", errors="ignore")
        return False, f"LINE API HTTP {e.code}: {err_body}"
    except Exception as e:
        return False, f"Error: {e}"

def get_libreoffice_cmd():
    for cmd in ["libreoffice", "soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"]:
        if shutil.which(cmd) or os.path.exists(cmd):
            return cmd
    return None

def process_and_convert(uploaded_file):
    lo_cmd = get_libreoffice_cmd()
    if not lo_cmd:
        raise RuntimeError("ไม่พบ LibreOffice ในระบบ (หากรันบน Streamlit Cloud ตรวจสอบว่ามีไฟล์ packages.txt ที่มีคำว่า libreoffice)")

    with tempfile.TemporaryDirectory() as tmpdir:
        input_path = os.path.join(tmpdir, "input.xlsx")
        with open(input_path, "wb") as f:
            f.write(uploaded_file.getbuffer())

        # ดึงยอดสรุปจากไฟล์
        summary = extract_summary_data(input_path)

        # ==========================================
        # ชีท 1: ยอดค้างรับลูกค้า
        # ==========================================
        wb1 = openpyxl.load_workbook(input_path)
        pdf_bytes_1 = None
        if "ยอดค้างรับลูกค้า" in wb1.sheetnames:
            ws1 = wb1["ยอดค้างรับลูกค้า"]
            for name in list(wb1.sheetnames):
                if name != "ยอดค้างรับลูกค้า":
                    wb1.remove(wb1[name])

            # ค้นหาแถวสิ้นสุด (ตารางสรุปยอดรวมวางบิลทั้งหมด)
            cutoff_row_1 = 925
            for r in range(1, ws1.max_row + 1):
                for c in range(1, 10):
                    val = str(ws1.cell(r, c).value or "")
                    if "ยอดรวมวางบิลทั้งหมด" in val:
                        cutoff_row_1 = r
                        break
                if cutoff_row_1 != 925:
                    break

            # กำหนดขอบเขตพิมพ์ A1:H{cutoff_row_1} และตั้งค่าหน้ากระดาษ A4 แนวตั้ง
            ws1.print_area = f"A1:H{cutoff_row_1}"
            ws1.page_setup.orientation = ws1.ORIENTATION_PORTRAIT
            ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
            ws1.page_setup.fitToWidth = 1
            ws1.page_setup.fitToHeight = 0
            ws1.sheet_properties.pageSetUpPr.fitToPage = True

            sheet1_xlsx = os.path.join(tmpdir, "sheet1.xlsx")
            wb1.save(sheet1_xlsx)

            subprocess.run([
                lo_cmd, "--headless", "--convert-to", "pdf",
                "--outdir", tmpdir, sheet1_xlsx
            ], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)

            pdf1_path = os.path.join(tmpdir, "sheet1.pdf")
            if os.path.exists(pdf1_path):
                with open(pdf1_path, "rb") as f:
                    pdf_bytes_1 = f.read()

        # ==========================================
        # ชีท 2: ยอดค้างจ่ายหัวลาก
        # ==========================================
        wb2 = openpyxl.load_workbook(input_path)
        pdf_bytes_2 = None
        if "ยอดค้างจ่ายหัวลาก" in wb2.sheetnames:
            ws2 = wb2["ยอดค้างจ่ายหัวลาก"]
            for name in list(wb2.sheetnames):
                if name != "ยอดค้างจ่ายหัวลาก":
                    wb2.remove(wb2[name])

            # ค้นหาแถวสิ้นสุด (หลังแถว Update สรุปยอดรวม)
            cutoff_row_2 = 92
            for r in range(1, ws2.max_row + 1):
                for c in range(1, 10):
                    val = str(ws2.cell(r, c).value or "")
                    if "Update" in val:
                        cutoff_row_2 = r + 2
                        break
                if cutoff_row_2 != 92:
                    break

            # กำหนดขอบเขตพิมพ์ A1:F{cutoff_row_2} และตั้งค่าหน้ากระดาษ A4 แนวตั้ง
            ws2.print_area = f"A1:F{cutoff_row_2}"
            ws2.page_setup.orientation = ws2.ORIENTATION_PORTRAIT
            ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
            ws2.page_setup.fitToWidth = 1
            ws2.page_setup.fitToHeight = 0
            ws2.sheet_properties.pageSetUpPr.fitToPage = True

            sheet2_xlsx = os.path.join(tmpdir, "sheet2.xlsx")
            wb2.save(sheet2_xlsx)

            subprocess.run([
                lo_cmd, "--headless", "--convert-to", "pdf",
                "--outdir", tmpdir, sheet2_xlsx
            ], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)

            pdf2_path = os.path.join(tmpdir, "sheet2.pdf")
            if os.path.exists(pdf2_path):
                with open(pdf2_path, "rb") as f:
                    pdf_bytes_2 = f.read()

        # ==========================================
        # รวมไฟล์ PDF (Combined) ถ้ามี pypdf ติดตั้ง
        # ==========================================
        combined_bytes = None
        if HAS_PYPDF and pdf_bytes_1 and pdf_bytes_2:
            merger = PdfMerger()
            merger.append(os.path.join(tmpdir, "sheet1.pdf"))
            merger.append(os.path.join(tmpdir, "sheet2.pdf"))
            comb_path = os.path.join(tmpdir, "combined.pdf")
            merger.write(comb_path)
            merger.close()
            with open(comb_path, "rb") as f:
                combined_bytes = f.read()

        return pdf_bytes_1, pdf_bytes_2, combined_bytes, summary

uploaded_file = st.file_uploader("ลากไฟล์ Excel (.xlsx) มาวางที่นี่", type=["xlsx"])

if uploaded_file is not None:
    with st.spinner("กำลังประมวลผลและสร้างไฟล์ PDF กรุณารอสักครู่..."):
        try:
            pdf1, pdf2, pdf_comb, summary = process_and_convert(uploaded_file)
            
            # ส่งการแจ้งเตือนเข้า LINE พร้อมยอดสรุป
            line_ok, line_msg = send_line_notification(uploaded_file.name, summary)
            
            st.success("✅ สร้างไฟล์ PDF สำเร็จเรียบร้อยแล้ว!")
            
            # แสดงกล่องสรุปยอดบนหน้าเว็บ
            st.write("---")
            st.subheader("📊 สรุปยอดจากเอกสาร:")
            col_m1, col_m2 = st.columns(2)
            with col_m1:
                total_rcv = f"{summary['ws1_total']:,.2f} บาท" if summary.get("ws1_total") else "-"
                st.metric(label="ยอดค้างรับทั้งหมด", value=total_rcv, help=summary.get("ws1_update"))
            with col_m2:
                total_pay = f"{summary['ws2_total']:,.2f} บาท" if summary.get("ws2_total") else "-"
                st.metric(label="ยอดค้างจ่ายหัวลาก", value=total_pay, help=summary.get("ws2_update"))

            # แสดงสถานะ LINE เพื่อให้ตรวจสอบได้ง่าย
            if line_ok:
                st.toast("📲 ส่งแจ้งเตือนเข้า LINE เรียบร้อยแล้ว!", icon="✅")
            else:
                st.info(f"ℹ️ สถานะ LINE แจ้งเตือน: {line_msg}")

            st.write("---")
            st.subheader("เลือกลิงก์ดาวน์โหลดไฟล์:")

            col1, col2 = st.columns(2)
            with col1:
                if pdf1:
                    st.download_button(
                        label="📥 ดาวน์โหลด ยอดค้างรับลูกค้า (PDF)",
                        data=pdf1,
                        file_name="ยอดค้างรับลูกค้า.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )
            with col2:
                if pdf2:
                    st.download_button(
                        label="📥 ดาวน์โหลด ยอดค้างจ่ายหัวลาก (PDF)",
                        data=pdf2,
                        file_name="ยอดค้างจ่ายหัวลาก.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )

            if pdf_comb:
                st.download_button(
                    label="📑 ดาวน์โหลด รวมทั้งสองส่วนในไฟล์เดียว (PDF)",
                    data=pdf_comb,
                    file_name="รายงานสรุปยอดค้างรับ-ค้างจ่าย.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )

        except Exception as e:
            st.error(f"❌ เกิดข้อผิดพลาด: {e}")
